#!/usr/bin/env python3
"""Reproducible, fictional CVs in genuine file formats, created offline.

The simulated desk-camera images are NOT actual phone photographs.
"""
import argparse
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont, ImageFilter, features
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

CV={
'en':('Amina Ahmad','Education','Bachelor of Education','Work Experience',
      'Teacher at Kabul school','Skills','Teaching and Microsoft Excel','Languages','English and Pashto'),
'fa':('فاطمه احمد','تحصیلات','لیسانس دانشگاه کابل','تجربه کاری',
      'آموزگار در مکتب','مهارت ها','تدریس و اکسل','زبان ها','فارسی و انگلیسی'),
'ps':('احمد خان','زده کړې','لېسانس د کابل پوهنتون','کاري تجربه',
      'ښوونکی په ښوونځي کې','مهارتونه','ښوونه او اکسل','ژبې','پښتو او انګلیسي')
}
def image_cv(lang,lines):
 im=Image.new('RGB',(1480,2020),'#fdfcf9')
 draw=ImageDraw.Draw(im)
 root='/usr/share/fonts/truetype/dejavu/DejaVuSans' if lang=='en' else '/usr/share/fonts/truetype/noto/NotoNaskhArabic'
 regular=ImageFont.truetype(root+('.ttf' if lang=='en' else '-Regular.ttf'),60)
 bold=ImageFont.truetype(root+'-Bold.ttf',67)
 draw.rectangle((90,85,1390,1920),outline='#c3c4c4',width=2)
 draw.line((140,260,1340,260),fill='#959c9e',width=3)
 for i,s in enumerate(lines):
  params={'font':bold if i in (0,1,3,5,7) else regular,'fill':'#14262f'}
  if lang=='en':draw.text((155,130+i*173),s,**params)
  else:draw.text((1320,130+i*173),s,anchor='ra',direction='rtl',**params)
 return im
def camera_simulation(im):
 base=Image.new('RGBA',(1700,2350),'#9a9b93')
 sheet=im.rotate(-1.8,resample=Image.Resampling.BICUBIC,expand=True,fillcolor='#ede9df')
 sheet.thumbnail((1450,2110))
 shadow=Image.new('RGBA',base.size,(0,0,0,0))
 ImageDraw.Draw(shadow).rounded_rectangle((154,115,1570,2250),radius=28,fill=(0,0,0,90))
 base=Image.alpha_composite(base,shadow.filter(ImageFilter.GaussianBlur(33)))
 base.alpha_composite(sheet.convert('RGBA'),((1700-sheet.width)//2,145))
 return base.convert('RGB').filter(ImageFilter.GaussianBlur(.3))
def create_docx(dest,lang,lines):
 doc=Document()
 for i,s in enumerate(lines):
  p=doc.add_paragraph()
  if lang!='en':
   from docx.enum.text import WD_ALIGN_PARAGRAPH
   p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  if i==0:p.add_run(s[:4]);p.add_run(s[4:])
  else:p.add_run(s)
 doc.save(dest)
def create_pdf(dest,lines):
 doc=Canvas(str(dest),pagesize=A4)
 for i,s in enumerate(lines):
  doc.setFont('Helvetica-Bold' if i in (0,1,3,5,7) else 'Helvetica',15 if i==0 else 12)
  doc.drawString(70,780-i*76,s)
 doc.save()
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--out',required=True)
 out=Path(ap.parse_args().out);out.mkdir(parents=True,exist_ok=True)
 if not features.check('raqm'):raise RuntimeError('libraqm missing for RTL rendering')
 for lang,lines in CV.items():
  create_docx(out/f'{lang}-native.docx',lang,lines)
  rendered=image_cv(lang,lines)
  rendered.save(out/f'{lang}-scanned.pdf','PDF',resolution=150.0)
  picture=camera_simulation(rendered)
  picture.save(out/f'{lang}-photo.jpg',quality=89,subsampling=0)
  if lang=='en':
   create_pdf(out/'en-native.pdf',lines)
   picture.save(out/'en-photo.webp',quality=86)
   picture.save(out/'en-photo.png')
 for path in out.iterdir():
  if path.stat().st_size>8*1024*1024:raise ValueError('oversized fixture')
 print('CV fixtures:',len(list(out.iterdir())),'real-format fictional documents; images simulate desk photos')
if __name__=='__main__':main()
